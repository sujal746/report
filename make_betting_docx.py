#!/usr/bin/env python3
"""Generate a premium professional DOCX: 'The Betting Process — Standard Operating Guide'."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageChops
import os
import tempfile

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
GRAY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x26, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SCREENSHOTS = "/home/user/report/screenshots"

doc = Document()

# ---------------- core properties ----------------
cp = doc.core_properties
cp.title = "The Betting Process — Standard Operating Guide"
cp.subject = "End-to-end workflow for placing, tracking, and settling sports bets"
cp.author = "Sujal"
cp.category = "Standard Operating Procedure"
cp.keywords = "betting process, sportsbook, bankroll, bet slip, cash out, responsible gambling"
cp.comments = "Premium professional betting process document, v1.0 — 23 October 2026"

# ---------------- base styles ----------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

for name, size, color, before, after in [
    ("Heading 1", 17, NAVY, 18, 8),
    ("Heading 2", 13.5, BLUE, 12, 6),
    ("Heading 3", 11.5, NAVY, 10, 4),
]:
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# ---------------- helpers ----------------
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=90, bottom=90, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)

def format_table(table):
    """Use fixed columns and subtle borders for a clean professional layout."""
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D5DDE8")
        borders.append(element)
    tblPr.append(borders)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trPr.append(no_split)

def add_bottom_border(paragraph, color="1F3864", sz="12", space="4"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def para(text="", size=11, bold=False, italic=False, color=None, align=None,
         space_before=None, space_after=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return p

def rich_para(chunks, align=None, space_after=None, style=None):
    """chunks: list of (text, dict) tuples."""
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    for text, fmt in chunks:
        r = p.add_run(text)
        r.font.size = Pt(fmt.get("size", 11))
        r.bold = fmt.get("bold", False)
        r.italic = fmt.get("italic", False)
        if "color" in fmt:
            r.font.color.rgb = fmt["color"]
    return p

def bullets(items, size=11):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            lead, rest = it
            r = p.add_run(lead)
            r.bold = True
            r.font.size = Pt(size)
            r2 = p.add_run(rest)
            r2.font.size = Pt(size)
        else:
            r = p.add_run(it)
            r.font.size = Pt(size)

def numbered(items, size=11):
    """Add a compact numbered list that restarts at 1 for each process block."""
    for index, it in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(4)
        n = p.add_run(f"{index}.  ")
        n.bold = True
        n.font.size = Pt(size)
        n.font.color.rgb = BLUE
        if isinstance(it, tuple):
            lead, rest = it
            r = p.add_run(lead)
            r.bold = True
            r.font.size = Pt(size)
            r2 = p.add_run(rest)
            r2.font.size = Pt(size)
        else:
            r = p.add_run(it)
            r.font.size = Pt(size)

def add_picture_centered(path, width_inches, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width_inches))
    if caption:
        para(caption, size=9, italic=True, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    return p

def callout(label, text, kind="info"):
    fills = {"info": ("EEF3FA", "1F3864"), "warning": ("FDF3E3", "B45309"),
             "success": ("EAF7EE", "15803D")}
    fill, bar = fills[kind]
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = doc.styles["Table Grid"]
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "28"); left.set(qn("w:color"), bar)
    borders.append(left)
    for side in ("top", "right", "bottom"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), bar)
        borders.append(el)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + "  ")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(bar)
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl

def info_table(rows, label_width=1.5):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = doc.styles["Table Grid"]
    tbl.autofit = False
    tbl.columns[0].width = Inches(label_width)
    tbl.columns[1].width = Inches(6.5 - label_width)
    for label, value in rows:
        row = tbl.add_row().cells
        set_cell_shading(row[0], "EEF3FA")
        set_cell_margins(row[0]); set_cell_margins(row[1])
        row[0].width = Inches(label_width)
        row[1].width = Inches(6.5 - label_width)
        p0 = row[0].paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = NAVY
        p1 = row[1].paragraphs[0]
        r1 = p1.add_run(value)
        r1.font.size = Pt(10)
    format_table(tbl)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl

def data_table(headers, rows, widths, header_fill="1F3864", size=9.5):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = doc.styles["Table Grid"]
    tbl.autofit = False
    for i, width in enumerate(widths):
        tbl.columns[i].width = Inches(width)
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].width = Inches(widths[i])
        set_cell_shading(hdr[i], header_fill)
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = WHITE
    header_props = tbl.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_props.append(repeat_header)
    for row_index, row in enumerate(rows, start=1):
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_margins(cells[i])
            if row_index % 2 == 0:
                set_cell_shading(cells[i], "F5F8FC")
            p = cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(size)
    format_table(tbl)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl

def field(paragraph, instr):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    return run

ASSET_TEMP = tempfile.TemporaryDirectory(prefix="betting-guide-")

def trim_white_border(path):
    """Trim empty white margins for legibility; never alter the source image."""
    image = Image.open(path).convert("RGB")
    diff = ImageChops.difference(image, Image.new("RGB", image.size, "white")).convert("L")
    mask = diff.point(lambda value: 255 if value > 18 else 0)
    bbox = mask.getbbox()
    if not bbox:
        return path
    pad = 10
    left = max(0, bbox[0] - pad)
    top = max(0, bbox[1] - pad)
    right = min(image.width, bbox[2] + pad)
    bottom = min(image.height, bbox[3] + pad)
    cropped = image.crop((left, top, right, bottom))
    output = os.path.join(ASSET_TEMP.name, "trimmed_" + os.path.basename(path))
    cropped.save(output, quality=95)
    return output

FIGURE = [0]
def figure(path, width, caption):
    FIGURE[0] += 1
    add_picture_centered(path, width, caption=f"Figure {FIGURE[0]} — {caption}")

# ---------------- page setup ----------------
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.9)
section.bottom_margin = Inches(0.9)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.different_first_page_header_footer = True

# header (from page 2 onward)
hp = section.header.paragraphs[0]
hp.text = ""
hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
r = hp.add_run("The Betting Process — Standard Operating Guide")
r.font.size = Pt(8.5); r.font.color.rgb = GRAY; r.bold = True
r2 = hp.add_run("\tProfessional Process Guide")
r2.font.size = Pt(8.5); r2.font.color.rgb = GRAY
add_bottom_border(hp, color="9AA5B1", sz="4", space="2")

# footer (from page 2 onward)
fp = section.footer.paragraphs[0]
fp.text = ""
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Prepared by Sujal  |  v1.0  |  Page ")
r.font.size = Pt(8.5); r.font.color.rgb = GRAY
field(fp, "PAGE")
r = fp.add_run(" of ")
r.font.size = Pt(8.5); r.font.color.rgb = GRAY
field(fp, "NUMPAGES")
for run in fp.runs:
    run.font.size = Pt(8.5); run.font.color.rgb = GRAY

# ---------------- cover page ----------------
para(space_after=2)
para("THE BETTING PROCESS", size=34, bold=True, color=NAVY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=60, space_after=4)
para("Standard Operating Guide for Placing, Tracking, and Settling Sports Bets",
     size=14, italic=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Operational workflow  •  Practical controls  •  Platform reference", size=11, color=GRAY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)

hp_cover = para(space_after=26)
add_bottom_border(hp_cover, color="2E74B5", sz="16", space="6")

info_table([
    ("Document title", "The Betting Process — Standard Operating Guide"),
    ("Prepared by", "Sujal"),
    ("Date", "23 October 2026"),
    ("Version", "1.0 (Issued)"),
    ("Applies to", "Stake · BC.Game · Roobet · Rollbit · Shuffle · Gamdom"),
    ("Document type", "Professional process reference"),
], label_width=1.7)

para(space_after=30)
callout("Purpose.",
        "This document defines the complete, step-by-step betting process — from account setup and "
        "bankroll management through bet placement, live tracking, settlement, and withdrawal. It is "
        "intended as a reference for adults placing sports bets with online bookmakers where legally permitted.",
        kind="info")
callout("Responsible gambling.",
        "Betting is entertainment, not a source of income. Only bet what you can afford to lose, set "
        "deposit and loss limits before you start, and never chase losses. If betting stops being fun, "
        "stop — and use the self-exclusion tools or support resources listed in Section 2.",
        kind="warning")

doc.add_page_break()

# ---------------- document control & contents ----------------
doc.add_heading("Document Control", level=1)
data_table(
    ["Version", "Date", "Author", "Change summary", "Status"],
    [["1.0", "23 Oct 2026", "Sujal", "Initial release — end-to-end betting process with platform-specific notes", "Issued"]],
    [0.8, 1.0, 1.0, 2.6, 1.1],
)

para(space_after=4)
doc.add_heading("Contents", level=1)
toc_items = [
    "1.  Purpose and Scope",
    "2.  Responsible Gambling — Read First",
    "3.  Glossary of Betting Terms",
    "4.  Phase 1 — Account and Bankroll Setup",
    "5.  Phase 2 — Placing a Bet",
    "6.  Phase 3 — During the Event",
    "7.  Phase 4 — Settlement and Withdrawal",
    "8.  Bet Types Quick Reference",
    "9.  Platform Playbook (Screenshot Notes)",
    "10.  Appendix — Screenshot Index",
]
for t in toc_items:
    para(t, size=11, color=NAVY, space_after=3)

doc.add_page_break()

# ---------------- 1. purpose ----------------
doc.add_heading("1.  Purpose and Scope", level=1)
para("This guide standardises how a sports bet is placed, from account setup until any winnings are "
     "withdrawn. It draws on real mobile and desktop product screens from Stake, BC.Game, Roobet, Rollbit, "
     "Shuffle and Gamdom. The screenshots illustrate navigation and review points; they are not endorsements, "
     "and interfaces, features and availability may change.")
para("Scope:", bold=True, space_after=2)
bullets([
    ("Account registration, ", "identity verification (KYC) and deposit."),
    ("Choosing sports, markets ", "and reading odds."),
    ("Building the bet slip, ", "setting stake, and confirming the bet."),
    ("Tracking live bets, ", "cash out, and in-play markets."),
    ("Settlement check ", "and withdrawal of winnings."),
])
para("Out of scope:", bold=True, space_after=2)
bullets([
    "Casino games, slots, crash games and raffles (except where noted).",
    "Financial or investment advice — betting outcomes are never guaranteed.",
    "Legal advice — confirm that online betting is permitted in your jurisdiction before you start.",
])

# ---------------- 2. responsible gambling ----------------
doc.add_heading("2.  Responsible Gambling — Read First", level=1)
callout("Golden rules.",
        "1. Participate only if you meet the legal age and betting is permitted where you live.  "
        "2. Treat every stake as money you may lose.  3. Set a bankroll and never exceed it.  "
        "4. Set deposit, loss and time limits before your first bet.  5. Never chase losses or borrow to bet.  "
        "6. If betting affects your mood, sleep or finances, stop and seek support.", kind="warning")
para("Platform tools to use:", bold=True, space_after=2)
bullets([
    "Deposit limits (daily / weekly / monthly).",
    "Loss limits and session time limits.",
    "Reality checks (session duration reminders).",
    "Self-exclusion / cool-off periods.",
    "Two-factor authentication (2FA) to secure the account.",
])
para("Support: Gamblers Anonymous (gamblersanonymous.org) or a qualified local gambling-support service. "
     "For legal questions, check the rules that apply in your state or country before placing a real-money bet.",
     size=10, italic=True, color=GRAY)

# ---------------- 3. glossary ----------------
doc.add_heading("3.  Glossary of Betting Terms", level=1)
data_table(
    ["Term", "Meaning"],
    [
        ["Stake", "The amount of money risked on a bet."],
        ["Odds (decimal)", "The multiplier applied to the stake to calculate the total return (stake × odds). Example: ₹1,000 @ 2.50 returns ₹2,500."],
        ["Market", "The specific outcome being bet on (e.g. match winner, over/under goals, both teams to score)."],
        ["Selection", "One outcome added to the bet slip."],
        ["Bet slip", "The panel where selections, stake, and potential return are shown before confirming."],
        ["Parlay / Accumulator", "Multiple selections combined into one bet; all must win for the bet to pay. Higher return, higher risk."],
        ["Cash out", "Settling a live bet early for an offered amount before the event ends."],
        ["In-play / Live", "Betting on an event after it has started, with odds that move in real time."],
        ["KYC", "Know Your Customer — identity checks that a platform may require before deposits or withdrawals."],
        ["Bankroll", "The total amount of money set aside exclusively for betting."],
        ["Settlement", "The bookmaker grading a bet as won, lost, or void after the event finishes."],
        ["RTP", "Return to Player — the theoretical percentage returned to players over time (casino products)."],
    ],
    [1.5, 5.0],
)

doc.add_page_break()

# ---------------- 4. phase 1 ----------------
doc.add_heading("4.  Phase 1 — Account and Bankroll Setup", level=1)

doc.add_heading("Step 1 — Choose a platform", level=2)
para("Pick a bookmaker that is reputable, holds a recognisable licence, covers the sports you want, and "
     "offers a clean mobile experience. The table below summarises the six platforms reviewed for this "
     "guide, based on live product screenshots.")
data_table(
    ["Platform", "Device focus", "Standout feature (from screenshots)", "Best used for"],
    [
        ["Stake", "Mobile", "Live counters on sport chips; full 1×2 market priced inline in the list", "Live soccer and fast market navigation"],
        ["BC.Game", "Mobile", "Quick-bet button inside the odds row; big-win social proof above the fold", "Quick bet placement on mobile"],
        ["Roobet", "Mobile / Desktop", "My Bets as a top-level sports tab; per-sport in-play state line", "Tracking settled and active bets"],
        ["Rollbit", "Desktop", "Combo quick-bet builder; promotional bets feed on the home page", "Parlays and combo bets"],
        ["Shuffle", "Mobile / Desktop", "Clean tennis bet slip on mobile; dense desktop sportsbook", "Tennis and general sportsbook browsing"],
        ["Gamdom", "Desktop", "Sports hub layout; rewards tiers visible in the account", "Bonus tracking and rewards status"],
    ],
    [1.0, 1.0, 2.6, 1.9],
)

doc.add_heading("Step 2 — Register and verify the account", level=2)
numbered([
    ("Open the registration form ", "and enter accurate email, phone and date-of-birth details. Use a strong, unique password."),
    ("Enable two-factor authentication ", "where available to protect the account and balance."),
    ("Confirm your email or phone ", "using the one-time code. The BC.Game example masks the destination and provides "
    "Paste and resend controls; email confirmation is separate from identity verification."),
    ("Complete KYC when requested ", "using the platform's secure upload flow. Requirements vary; check withdrawal "
    "requirements before depositing significant funds."),
])
figure(f"{SCREENSHOTS}/05_bcgame_mobile_email_verification.jpg", 2.4,
       "BC.Game mobile: email verification with a masked destination, Paste control and resend countdown.")

doc.add_heading("Step 3 — Deposit and set your bankroll", level=2)
numbered([
    ("Decide your bankroll ", "— money you can afford to lose, kept separate from essential and daily finances."),
    ("Deposit through the cashier ", "using an available method; check fees, payment network, processing time and any bonus conditions first."),
    ("Set a fixed unit size ", "before betting. An illustrative guardrail is 1–2% of bankroll per bet; this limits exposure but does not improve the odds."),
    ("Apply deposit, loss and time limits ", "in the account settings before the first bet (see Section 2)."),
])
info_table([
    ("Example guardrail", "1–2% of bankroll per bet; never increase a stake to chase losses."),
    ("Before depositing", "Check verification, withdrawal, fee and payment-method requirements."),
    ("Security", "Enable 2FA where available; never share passwords or one-time codes."),
])

doc.add_page_break()

# ---------------- 5. phase 2 ----------------
doc.add_heading("5.  Phase 2 — Placing a Bet", level=1)

doc.add_heading("Step 4 — Pick the sport and market", level=2)
para("Open the sportsbook, choose a sport, then a competition and event. Look for live indicators — "
     "Stake, for example, shows live counters on the sport chips (Tennis 48, Soccer 4, Baseball 8) with the "
     "current score sitting inside the same card as the price, so you can enter a live market in one tap.")
figure(f"{SCREENSHOTS}/01_stake_mobile_live_soccer.jpg", 2.5,
       "Stake mobile: live counters on sport chips and the score inside the event card.")

doc.add_heading("Step 5 — Read the odds and implied probability", level=2)
para("These examples use decimal odds. Implied probability is 1 ÷ odds — for example, decimal odds of "
     "2.00 convert to 50%. This is a mathematical conversion, not a prediction or guarantee. Stake's mobile "
     "list prices the full 1×2 row inline for each match, with a chip for additional markets, so you can "
     "review the event, market and price before opening the slip.")
figure(f"{SCREENSHOTS}/02_stake_mobile_market_density.jpg", 2.1,
       "Stake mobile: the 1×2 market priced inline, with a chip showing additional markets.")

doc.add_heading("Step 6 — Add selections to the bet slip", level=2)
para("Tap the odds of your chosen outcome to add it to the slip. On BC.Game mobile, the green bet-slip "
     "button sits in the market area; on Shuffle mobile, Bet Slip remains available in the bottom bar. "
     "Open the slip and review every selection before staking — a bet normally cannot be edited after confirmation.")
figure(trim_white_border(f"{SCREENSHOTS}/03_bcgame_mobile_quickbet_fab.jpg"), 2.7,
       "BC.Game mobile: a prominent bet-slip control within the sports market view.")
figure(trim_white_border(f"{SCREENSHOTS}/12_shuffle_mobile_tennis_betslip.jpg"), 2.5,
       "Shuffle mobile: live tennis prices with Bet Slip retained as a bottom-bar destination.")

doc.add_heading("Step 7 — Set the stake and verify the potential return", level=2)
numbered([
    ("Enter your stake ", "within the fixed unit limit you set in Step 3."),
    ("Check the potential total return ", "— stake × decimal odds for a single, or stake × combined decimal odds for a parlay. Potential profit is total return minus stake."),
    ("Confirm the bet type ", "— single, parlay/accumulator or system (see Section 8)."),
    ("Review the final slip ", "— event, market, selection, odds, stake, return and any promotion conditions."),
])

doc.add_heading("Step 8 — Place the bet and confirm", level=2)
para("Tap “Place Bet” (or “Confirm”) and wait for an acceptance message or bet ID. If Quick Bet is enabled, "
     "check whether it reduces confirmation steps before using it. The Rollbit example below shows why the "
     "bet-type tab, leg count and each selection must be reviewed before submission.")
figure(f"{SCREENSHOTS}/10_rollbit_desktop_quickbet_combo.jpg", 5.9,
       "Rollbit desktop: Single / Combo / System tabs, two listed legs and a visible Quick Bet toggle.")

callout("Pro tip.",
        "Save the accepted bet details (bet ID, selections, stake, odds and timestamp) for any later query. "
        "If the connection drops during submission, check My Bets before retrying so you do not place a duplicate.",
        kind="info")

doc.add_page_break()

# ---------------- 6. phase 3 ----------------
doc.add_heading("6.  Phase 3 — During the Event", level=1)

doc.add_heading("Step 9 — Track the bet in “My Bets”", level=2)
para("Open “My Bets” from the sports navigation to review active and settled wagers, their status and "
     "potential returns. Roobet places My Bets beside All Sports in the top-level navigation. Use the bet "
     "ID and acceptance status—not a selection badge alone—as evidence that the wager was recorded.")
figure(f"{SCREENSHOTS}/06_roobet_mobile_nav_mybets.jpg", 2.5,
       "Roobet mobile: My Bets available as a first-class tab in the sports navigation.")

doc.add_heading("Step 10 — Live betting and cash out (optional)", level=2)
para("If the event is underway, use the in-play list, which shows each league's state line (halftime, "
     "minute, quarter) beside the market and price. Roobet's mobile in-play view names the market above "
     "every price, so you know exactly what you are backing at each live moment.")
figure(f"{SCREENSHOTS}/07_roobet_mobile_inplay.jpg", 2.5,
       "Roobet mobile: in-play state lines, scores and market labels shown alongside live prices.")
callout("Cash out — where offered.",
        "Cash out can settle an eligible live bet early for the amount shown. The offer moves with live "
        "odds and may be suspended or unavailable. Accepting it ends the bet at that value; you give up "
        "the original full payout if the selection later wins.", kind="info")

doc.add_page_break()

# ---------------- 7. phase 4 ----------------
doc.add_heading("7.  Phase 4 — Settlement and Withdrawal", level=1)

doc.add_heading("Step 11 — Check settlement", level=2)
numbered([
    ("Wait for the event to finish ", "— the bookmaker grades the bet as won, lost, or void."),
    ("Check “My Bets” / settled bets ", "for the result and the payout credited to your balance."),
    ("Query discrepancies quickly ", "using the bet ID and your screenshot from Step 8."),
])

doc.add_heading("Step 12 — Withdraw winnings", level=2)
numbered([
    ("Confirm KYC is complete ", "— most platforms block withdrawals from unverified accounts."),
    ("Open the cashier/withdrawal section ", "and choose your payout method and network."),
    ("Enter the amount ", "within min/max limits; note fees and expected processing time."),
    ("Withdraw to an account in your own name ", " — third-party payouts are usually rejected."),
])

callout("Withdrawal checklist.",
        "KYC complete ✅  •  2FA enabled ✅  •  Amount within limits ✅  •  Correct network/address ✅  •  "
        "Screenshot of withdrawal request ✅", kind="success")

# ---------------- 8. bet types ----------------
doc.add_heading("8.  Bet Types Quick Reference", level=1)
data_table(
    ["Bet type", "How it works", "Review priority", "Relative risk"],
    [
        ["Single", "One selection on one market.", "Confirm the event, market, selection and price.", "Lower than a parlay; stake still at risk"],
        ["Parlay / Accumulator", "Two or more selections; all must win.", "Check every leg and the combined odds.", "Higher; one losing leg loses the bet"],
        ["System", "A combination of smaller parlays, depending on the format.", "Check total cost, combination count and payout rules.", "Varies by format; not a guarantee"],
        ["In-play / Live", "Placed after the event starts; odds and availability move quickly.", "Check the score, game state, market label and acceptance.", "Fast-moving; suspension and price changes possible"],
    ],
    [1.25, 2.05, 2.05, 1.15],
)
para("Promotions may include trigger conditions, minimum stakes and maximum payouts. Read the terms before "
     "opting in. A promotion does not remove the risk of losing, and every added parlay leg adds another way "
     "for the bet to lose.", size=10, italic=True, color=GRAY)

doc.add_page_break()

# ---------------- 9. platform playbook ----------------
doc.add_heading("9.  Platform Playbook (Screenshot Notes)", level=1)
para("Use these observations as navigation shortcuts, not as rules or endorsements. Interfaces, market "
     "availability and feature names can change; always verify the live slip in front of you.")
data_table(
    ["Platform", "Workflow shortcut", "Illustrated in"],
    [
        ["Stake", "Use live counters to enter active events; compare inline 1×2 prices before opening a selection.", "Figures 2–3 · Steps 4–5"],
        ["BC.Game", "Complete account-code verification; use the sports-screen bet-slip control to review selections.", "Figures 1 and 4 · Steps 2 and 6"],
        ["Roobet", "Use My Bets for active and settled wagers; check game state and market labels on in-play events.", "Figures 7–8 · Steps 9–10"],
        ["Rollbit", "Choose Single, Combo or System explicitly; review every leg and the Quick Bet setting before submitting.", "Figure 6 · Step 8"],
        ["Shuffle", "Use the persistent Bet Slip destination in the bottom bar to return to your selections.", "Figure 5 · Step 6"],
        ["Gamdom", "Scan the sports hub for available markets; review rewards and promotional terms separately from bet rules.", "Reference screenshot pack"],
    ],
    [1.0, 4.0, 1.5],
)
callout("Interface check.",
        "A familiar icon or green/yellow highlight is not confirmation that a bet was accepted. Always "
        "verify the final status in the bet slip or My Bets and retain the bet ID where available.", kind="info")

doc.add_page_break()

# ---------------- 10. appendix ----------------
doc.add_heading("10.  Appendix — Screenshot Index", level=1)
para("The eight figures in this guide come from the real product screenshot pack in the project workspace. "
     "Two phone captures are trimmed only to remove blank white margins; their interface content is unchanged.",
     italic=True, size=9.5, color=GRAY, space_after=6)
data_table(
    ["Figure", "Platform / device", "Source file", "Process reference"],
    [
        ["1", "BC.Game / mobile", "05_bcgame_mobile_email_verification.jpg", "Step 2 · Verify account"],
        ["2", "Stake / mobile", "01_stake_mobile_live_soccer.jpg", "Step 4 · Select event"],
        ["3", "Stake / mobile", "02_stake_mobile_market_density.jpg", "Step 5 · Review odds"],
        ["4", "BC.Game / mobile", "03_bcgame_mobile_quickbet_fab.jpg", "Step 6 · Open bet slip"],
        ["5", "Shuffle / mobile", "12_shuffle_mobile_tennis_betslip.jpg", "Step 6 · Review selections"],
        ["6", "Rollbit / desktop", "10_rollbit_desktop_quickbet_combo.jpg", "Step 8 · Confirm bet"],
        ["7", "Roobet / mobile", "06_roobet_mobile_nav_mybets.jpg", "Step 9 · Track bet"],
        ["8", "Roobet / mobile", "07_roobet_mobile_inplay.jpg", "Step 10 · In-play review"],
    ],
    [0.65, 1.35, 3.25, 1.25],
    size=9,
)

para(space_after=10)
callout("End of document.",
        "Questions or process improvements should be sent to the document owner. This guide does not "
        "guarantee betting outcomes and is not financial or legal advice. Bet responsibly.", kind="info")

# Ask Word-compatible editors to refresh PAGE / NUMPAGES fields on open.
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
doc.settings.element.append(update_fields)

OUT = "/home/user/report/Betting_Process_Professional_Guide.docx"
doc.save(OUT)
print("Saved:", OUT)